from __future__ import annotations

from io import BytesIO
from pathlib import Path

from langchain_community.document_loaders.base import BaseLoader
from langchain_core.documents import Document
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .image_file_loader import ImageDocumentLoader


class PowerPointStyleLoader(BaseLoader):
    """
    PowerPoint(.pptx)用カスタムLoader

    1 Slide = 1 Document

    取得情報
    --------
    ・テキスト
    ・フォント
    ・文字色
    ・サイズ
    ・太字
    ・斜体
    ・下線
    ・段落レベル
    ・図形種類
    ・図形位置
    ・画像
    ・Speaker Notes
    """

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)

    def load(self) -> list[Document]:

        prs = Presentation(self.file_path)

        docs = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_parts = []

            ###########################################################
            # Shape
            ###########################################################

            for shape_idx, shape in enumerate(slide.shapes):
                ###################################################
                # Text
                ###################################################

                if shape.has_text_frame:
                    paragraphs = []

                    for paragraph in shape.text_frame.paragraphs:
                        runs = []

                        for run in paragraph.runs:
                            font = run.font

                            color = None
                            if (
                                font.color
                                and font.color.type
                                and hasattr(font.color, "rgb")
                                and font.color.rgb
                            ):
                                color = str(font.color.rgb)

                            runs.append(
                                {
                                    "text": run.text,
                                    "font_name": font.name,
                                    "font_size": (font.size.pt if font.size else None),
                                    "bold": font.bold,
                                    "italic": font.italic,
                                    "underline": font.underline,
                                    "font_color": color,
                                }
                            )

                        paragraphs.append(
                            {
                                "level": paragraph.level,
                                "runs": runs,
                            }
                        )

                    slide_parts.append(
                        {
                            "type": "text",
                            "data": {
                                "shape_index": shape_idx,
                                "shape_name": shape.name,
                                "shape_type": str(shape.shape_type),
                                "left": shape.left,
                                "top": shape.top,
                                "width": shape.width,
                                "height": shape.height,
                                "paragraphs": paragraphs,
                            },
                        }
                    )

                ###################################################
                # Table
                ###################################################

                if shape.has_table:
                    table_data = []

                    table = shape.table

                    for row in table.rows:
                        values = []

                        for cell in row.cells:
                            values.append(cell.text)

                        table_data.append(values)

                    slide_parts.append(
                        {
                            "type": "table",
                            "data": table_data,
                        }
                    )

                ###################################################
                # Picture
                ###################################################

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image = Image.open(BytesIO(shape.image.blob))

                    response_data, image_store_id = ImageDocumentLoader(
                        self.file_path
                    )._describe_image(image)

                    description = response_data[1]["text"]

                    slide_parts.append(
                        {
                            "type": "image",
                            "data": {
                                "shape_index": shape_idx,
                                "image_store_id": image_store_id,
                                "description": description,
                            },
                        }
                    )

            ###########################################################
            # Speaker Notes
            ###########################################################

            notes_text = None

            try:
                note_parts = []

                for shape in slide.notes_slide.shapes:
                    if hasattr(shape, "text"):
                        text = shape.text.strip()
                        if text:
                            note_parts.append(text)

                if note_parts:
                    notes_text = "\n".join(note_parts)

            except Exception:
                pass

            ###########################################################
            # metadata
            ###########################################################

            image_ids = [p["data"]["image_store_id"] for p in slide_parts if p["type"] == "image"]

            metadata = {
                "source": str(self.file_path),
                "slide_index": slide_idx,
                "layout": slide.slide_layout.name,
                "image_count": len(image_ids),
                "image_store_ids": image_ids if image_ids else None,
            }

            docs.append(
                Document(
                    page_content=self._build_text(
                        slide_idx,
                        slide_parts,
                        notes_text,
                    ),
                    metadata=metadata,
                )
            )

        return docs

    ##############################################################

    def _build_text(
        self,
        slide_index,
        slide_parts,
        notes_text,
    ):

        texts = []

        texts.append(f"# Slide {slide_index + 1}")
        texts.append("")

        for part in slide_parts:
            ######################################################
            # Text
            ######################################################

            if part["type"] == "text":
                shape = part["data"]

                texts.append(f"[TEXTBOX] {shape['shape_name']}")

                for paragraph in shape["paragraphs"]:
                    for run in paragraph["runs"]:
                        if not run["text"]:
                            continue

                        texts.append(run["text"])

                        if run["bold"]:
                            texts.append("[BOLD]")

                        if run["italic"]:
                            texts.append("[ITALIC]")

                        if run["underline"]:
                            texts.append("[UNDERLINE]")

                        if run["font_name"]:
                            texts.append(f"[FONT:{run['font_name']}]")

                        if run["font_size"]:
                            texts.append(f"[SIZE:{run['font_size']}]")

                        if run["font_color"]:
                            texts.append(f"[FONT_COLOR:{run['font_color']}]")

                texts.append("")

            ######################################################
            # Table
            ######################################################

            elif part["type"] == "table":
                texts.append("[TABLE]")

                for row in part["data"]:
                    texts.append(" | ".join(row))

                texts.append("")

            ######################################################
            # Image
            ######################################################

            elif part["type"] == "image":
                texts.append("[IMAGE]")
                texts.append(part["data"]["description"])
                texts.append("")

        if notes_text:
            texts.append("[SPEAKER NOTES]")
            texts.append(notes_text)

        return "\n".join(texts)


if __name__ == "__main__":
    # test
    from langchain_community.document_loaders import UnstructuredPowerPointLoader

    path = "../../share/共有ドライブ/プロジェクト/白峰信用リスク評価株式会社/06.報告書/白峰信用リスク評価株式会社_最終報告.pptx"
    docs = PowerPointStyleLoader(path).load()
    docs += UnstructuredPowerPointLoader(path).load()
    for doc in docs:
        print(doc.metadata)
        print(doc.page_content)
        print("--------------------------")
