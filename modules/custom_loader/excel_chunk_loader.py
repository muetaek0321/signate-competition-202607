from collections.abc import Iterator
from pathlib import Path

from langchain_community.document_loaders.base import BaseLoader
from langchain_core.documents import Document
from openpyxl import load_workbook


class ExcelChunkLoader(BaseLoader):
    """
    Excelをシート単位・行単位で分割するDocumentLoader

    Parameters
    ----------
    file_path : str | Path
        Excelファイルパス

    chunk_rows : int
        1 Documentあたりの行数

    include_header : bool
        各chunkにヘッダー行を含めるか

    header_rows : int
        ヘッダーとして扱う行数
    """

    def __init__(
        self,
        file_path: str | Path,
        chunk_rows: int = 300,
        include_header: bool = True,
        header_rows: int = 1,
    ):
        self.file_path = Path(file_path)
        self.chunk_rows = chunk_rows
        self.include_header = include_header
        self.header_rows = header_rows

    def load(self) -> list[Document]:
        return list(self.lazy_load())

    def lazy_load(self) -> Iterator[Document]:
        wb = load_workbook(
            filename=self.file_path,
            read_only=True,
            data_only=True,
        )

        for ws in wb.worksheets:
            max_row = ws.max_row
            max_col = ws.max_column

            # ヘッダー取得
            headers = []

            if self.include_header:
                for row in ws.iter_rows(
                    min_row=1,
                    max_row=self.header_rows,
                    values_only=True,
                ):
                    headers.append(list(row))

            start_row = self.header_rows + 1

            chunk_index = 0

            while start_row <= max_row:
                end_row = min(start_row + self.chunk_rows - 1, max_row)

                rows = []

                if self.include_header:
                    rows.extend(headers)

                for row in ws.iter_rows(
                    min_row=start_row,
                    max_row=end_row,
                    max_col=max_col,
                    values_only=True,
                ):
                    rows.append(list(row))

                text = self._rows_to_markdown(rows)

                yield Document(
                    page_content=text,
                    metadata={
                        "source": str(self.file_path),
                        "file_name": self.file_path.name,
                        "sheet_name": ws.title,
                        "chunk_index": chunk_index,
                        "start_row": start_row,
                        "end_row": end_row,
                        "total_rows": max_row,
                        "total_columns": max_col,
                    },
                )

                chunk_index += 1

                start_row = end_row + 1

        wb.close()

    def _rows_to_markdown(self, rows: list[list]) -> str:

        lines = []

        for row in rows:
            values = ["" if v is None else str(v) for v in row]

            lines.append("| " + " | ".join(values) + " |")

        return "\n".join(lines)
